# -*- coding: utf-8 -*-
"""
表格处理合规工具（配合 SKILL.md §四「表格处理铁律」）。
目的：让"按单元格读表 + 强制质检"成为默认动作，而不是靠记忆。

用法:
    python3 table_qa.py pdf   <file.pdf> [--page N] [--compare]   # 结构化读 PDF 表格 + 与扁平文本对照
    python3 table_qa.py docx  <file.docx>                         # 导出 docx 全部表格（行列号 + 内容）
    python3 table_qa.py pptx  <file.pptx>                         # 导出 PPTX 原生表格（a:tbl）逐格内容
    python3 table_qa.py sums  <a> <b> ...                         # 数字合计验算（含 % 自动归一提示）

要点：
- PDF 表格优先 find_tables()（有框线）；无框线退化为 text 策略；并给出表头 x 坐标映射，便于人工确认列归属。
- --compare 会把 extract_text() 的扁平文本一并打印：用于发现"文本流顺序 ≠ 单元格顺序"的错位陷阱。
- 任何脚本输出都必须与原始文件逐格抽查后才可写入笔记（见 SKILL §四 4.3）。
"""
import sys
import re


# ---------------------------------------------------------------- PDF
def pdf_tables(path, page=None, compare=False):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        pages = [pdf.pages[page - 1]] if page else pdf.pages
        for pi, p in enumerate(pages, start=(page or 1)):
            found = p.find_tables()
            print(f'===== PDF page {pi}: 检测到 {len(found)} 个表格 =====')
            if not found:
                # 无框线表格：用文本策略再试一次
                try:
                    tables = p.extract_tables(table_settings={
                        'vertical_strategy': 'text', 'horizontal_strategy': 'text'})
                    for ti, tb in enumerate(tables):
                        print(f'-- [text-strategy] table {ti + 1}: {len(tb)} 行 x {len(tb[0]) if tb else 0} 列')
                        _print_rows(tb)
                except Exception as e:
                    print('  [text-strategy 失败]', e)
            for ti, t in enumerate(found):
                rows = t.extract()
                ncol = max((len(r) for r in rows), default=0)
                print(f'-- table {ti + 1}: {len(rows)} 行 x {ncol} 列  bbox={tuple(round(v) for v in t.bbox)}')
                _print_rows(rows)
            # 表头 x 坐标映射（判断列归属的唯一可靠依据）
            words = p.extract_words()
            if words:
                heads = [(round(w['x0'], 1), w['text']) for w in words
                         if w['text'] in ('Time', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday',
                                          'Day', 'Date', 'Week', 'Topic', 'Assessment', 'Weight', 'Criterion',
                                          'Module', 'Workshop', 'Quizzes')]
                if heads:
                    print('-- 表头 x 坐标（列边界依据）:', heads[:12])
            if compare:
                print('-- [扁平 extract_text 对照] 注意：顺序 ≠ 单元格顺序，禁止据此判断列归属 --')
                print((p.extract_text() or '')[:1200])


def _print_rows(rows):
    for ri, row in enumerate(rows):
        cells = ['' if c is None else re.sub(r'\s+', ' ', str(c)).strip() for c in row]
        print(f'   r{ri}: {cells}')


# ---------------------------------------------------------------- DOCX
def docx_tables(path):
    from docx import Document

    d = Document(path)
    print(f'===== {path}: {len(d.tables)} 个表格 =====')
    for ti, t in enumerate(d.tables):
        ncol = len(t.columns)
        print(f'-- table {ti + 1}: {len(t.rows)} 行 x {ncol} 列')
        for ri, row in enumerate(t.rows):
            cells, seen = [], []
            for c in row.cells:                     # 合并单元格会重复返回同一对象
                key = id(c._tc)
                cells.append('⟦merged⟧' + c.text.strip() if key in seen else c.text.strip())
                seen.append(key)
            print(f'   r{ri}: {cells}')
    print('\n写完 docx 单元格后必须重新 dump 全表核对（empty/串列/重复）。')


def set_docx_cell(cell, text):
    """安全写单元格：清空既有 run 再写，避免文本重复。"""
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ''
    para = cell.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


# ---------------------------------------------------------------- PPTX
def pptx_tables(path):
    import zipfile
    from xml.etree import ElementTree as ET

    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
          'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    z = zipfile.ZipFile(path)
    slides = sorted([n for n in z.namelist()
                     if re.match(r'ppt/slides/slide\d+\.xml$', n)],
                    key=lambda s: int(re.findall(r'\d+', s)[0]))
    total = 0
    for s in slides:
        sid = int(re.findall(r'\d+', s)[0])
        root = ET.fromstring(z.read(s))
        for ti, tbl in enumerate(root.iter(f'{{{ns["a"]}}}tbl')):
            total += 1
            rows = list(tbl.findall(f'{{{ns["a"]}}}tr'))
            ncol = max((len(r.findall(f'{{{ns["a"]}}}tc')) for r in rows), default=0)
            print(f'-- slide {sid} table {ti + 1}: {len(rows)} 行 x {ncol} 列')
            for ri, tr in enumerate(rows):
                cells = []
                for tc in tr.findall(f'{{{ns["a"]}}}tc'):
                    txt = ''.join(t.text or '' for t in tc.iter(f'{{{ns["a"]}}}t'))
                    cells.append(re.sub(r'\s+', ' ', txt).strip())
                print(f'   r{ri}: {cells}')
            print('   ⚠ 合并单元格可能重复出现同一文本；列数取最大行单元格数')
    if total == 0:
        print('无原生表格；若课件有 SmartArt/图表，请检查 ppt/diagrams/data*.xml 与 ppt/charts/chart*.xml')


# ---------------------------------------------------------------- 合计验算
def sums(values):
    cleaned, pct = [], 0
    for v in values:
        s = v.strip()
        is_pct = s.endswith('%')
        num = re.sub(r'[^0-9.\-]', '', s)
        if num == '':
            continue
        f = float(num)
        cleaned.append(f)
        pct += 1 if is_pct else 0
    total = sum(cleaned)
    print(f'数值: {cleaned}')
    print(f'合计: {total}')
    if pct == len(cleaned) and cleaned:
        print('判定: ' + ('✅ 合计 = 100%（自洽）' if abs(total - 100) < 0.01
                        else f'⚠️ 合计 = {total} ≠ 100% → 照录原文并标"待老师确认"，不得自行改写'))
    else:
        print('判定: 非全百分比，请人工判断该合计应有的期望值。')


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return 1
    mode = args[0]
    if mode == 'pdf':
        page = None; compare = '--compare' in args
        if '--page' in args:
            page = int(args[args.index('--page') + 1])
        pdf_tables(args[1], page=page, compare=compare)
    elif mode == 'docx':
        docx_tables(args[1])
    elif mode == 'pptx':
        pptx_tables(args[1])
    elif mode == 'sums':
        sums(args[1:])
    else:
        print(__doc__)
        return 1
    return 0


try:  # Windows 控制台默认 cp1252/cp936，统一转 UTF-8，否则打印中文会 UnicodeEncodeError
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

if __name__ == '__main__':
    sys.exit(main())
